"""
md_to_docx.py — Universal Markdown → DOCX exporter for VNR deliverables.

Clones a branded DOCX template and rebuilds the body from a Markdown source,
preserving the template's styles, page setup, margins, header/footer and cover
page. Handles headings (auto-numbered via Word styles), paragraphs, bullet &
numbered lists, tables (auto column widths + captions), blockquotes, fenced
code blocks and inline bold/italic/`code`.

This is the generalized engine distilled from scripts/export_all.py — it takes
ANY markdown file instead of a hardcoded registry.

Usage:
    python md_to_docx.py INPUT.md [-o OUTPUT.docx|OUTDIR] [options]

Options:
    -o, --output PATH     Output .docx file, or a directory (filename derived
                          from the input). Default: alongside the input file.
    -t, --template PATH   DOCX template to clone. Default: the VNR-ARCH template.
    --code CODE           Document code (e.g. VNR-KH-001). Default: parsed from
                          the filename, else from metadata.
    --title TITLE         Cover/title override. Default: first H1 in the markdown.
    --version VER         Document version (e.g. v1.0). Default: from metadata.
    --project-code CODE   Project code shown on the cover. Default: <Mã dự án>.
    --date "TPHCM, MM/YYYY"  Cover date line. Default: current month/year.
    --author NAME         Author for the revision-history table. Default: metadata
                          "Author", else "SA Team".
    --no-cover            Skip the cover page + front matter (revision/approval/TOC).
    --quiet               Suppress per-file log lines.

Examples:
    python md_to_docx.py ../vnr.hrm.playbook/07-customer/VNR-KH-001-....md -o customer/deployment
    python md_to_docx.py report.md --code VNR-ROAD-001 --title "Technical Roadmap 2026"

Requires: python-docx  (pip install python-docx)
"""
import argparse
import datetime
import os
import re
import shutil
import sys

try:
    from docx import Document
    from docx.shared import Pt, RGBColor, Cm, Twips
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn, nsdecls
    from docx.oxml import parse_xml
except ImportError:
    sys.stderr.write(
        "ERROR: the 'python-docx' package is required.\n"
        "       Install it with:  pip install python-docx\n")
    sys.exit(2)

# ── Locate the VNR-ARCH template + logo. When run without an explicit --template
#    (e.g. the export-docx wrappers / raw python), search the current working dir
#    (the deliverables repo the user runs in) first, then the script dir. Matches
#    both the repo-local `templates/word/` layout and the VNR.Platform superrepo
#    `src/vnr.hrm.deliverables/templates/word/` layout. The Claude Code skill passes
#    --template explicitly, so it never depends on this lookup. ──
_SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
_REL_LAYOUTS = (
    os.path.join("templates", "word"),                                 # repo-local (deliverables repo)
    os.path.join("src", "vnr.hrm.deliverables", "templates", "word"),  # VNR.Platform superrepo
)
_TEMPLATE_NAME = "VNR-ARCH_Tai_Lieu_Dac_Ta_Kien_Truc_He_Thong.docx"


def _find_template_dir():
    # Search CWD (the repo the user runs in) first, then the script dir.
    for start in (os.getcwd(), _SCRIPT_DIR):
        d = start
        for _ in range(8):  # walk up a bounded number of levels
            for rel in _REL_LAYOUTS:
                cand = os.path.join(d, rel)
                if os.path.isfile(os.path.join(cand, _TEMPLATE_NAME)):
                    return cand
            parent = os.path.dirname(d)
            if parent == d:
                break
            d = parent
    # Fallback: superrepo layout, 3 levels up from the script.
    return os.path.normpath(os.path.join(_SCRIPT_DIR, "..", "..", "..", _REL_LAYOUTS[1]))


_TEMPLATE_DIR = _find_template_dir()
DEFAULT_TEMPLATE = os.path.join(_TEMPLATE_DIR, _TEMPLATE_NAME)
DEFAULT_LOGO = os.path.join(_TEMPLATE_DIR, "logo-vnresource.png")

# Brand palette (matches the template).
C_TEXT = '1A1A1A'
C_BRAND = '1F4E79'
C_ACCENT = '2E75B6'
C_MUTED = '595959'
C_HEADER_FILL = '1F4E79'
C_HEADER_TEXT = 'FFFFFF'


# ──────────────────────────────────────────────
# DOCX HELPERS (match template styles exactly)
# ──────────────────────────────────────────────

def clear_body(doc):
    body = doc.element.body
    for child in list(body):
        if child.tag != qn('w:sectPr'):
            body.remove(child)


def _empty(doc):
    return doc.add_paragraph(style='Normal')


def _h1_no_num(doc, text):
    p = doc.add_paragraph(style='Heading 1 - No Numeric')
    for r in p.runs:
        r.text = ''
    p.add_run(text)
    return p


def _section_break(doc):
    """Add a section break (new page) before the next paragraph."""
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    sectPr = parse_xml(
        f'<w:sectPr {nsdecls("w")}>'
        f'<w:type w:val="nextPage"/>'
        f'</w:sectPr>'
    )
    pPr.append(sectPr)
    return p


def _h1(doc, text):
    text = text.replace('`', '')
    _section_break(doc)
    return doc.add_heading(text, level=1)


def _h2(doc, text):
    return doc.add_heading(text.replace('`', ''), level=2)


def _h3(doc, text):
    return doc.add_heading(text.replace('`', ''), level=3)


def _normal(doc, text, bold=False):
    p = doc.add_paragraph(style='Normal')
    r = p.add_run(text)
    r.font.bold = bold
    r.font.color.rgb = RGBColor.from_string(C_TEXT)
    return p


def _apply_runs(p, parts):
    """parts = [(text, bold, italic, code?), ...]"""
    for item in parts:
        text, bold, italic = item[0], item[1], item[2]
        is_code = item[3] if len(item) > 3 else False
        r = p.add_run(text)
        r.font.color.rgb = RGBColor.from_string(C_TEXT)
        if bold:
            r.font.bold = True
        if italic:
            r.font.italic = True
        if is_code:
            r.font.name = 'Courier New'
            r.font.size = Pt(9)
    return p


def _normal_rich(doc, parts):
    return _apply_runs(doc.add_paragraph(style='Normal'), parts)


def _set_bullet_numPr(p, num_id=2, ilvl=0):
    pPr = p._p.get_or_add_pPr()
    numPr = parse_xml(
        f'<w:numPr {nsdecls("w")}>'
        f'<w:ilvl w:val="{ilvl}"/>'
        f'<w:numId w:val="{num_id}"/>'
        f'</w:numPr>'
    )
    pPr.append(numPr)


def _bullet_rich(doc, parts):
    p = doc.add_paragraph(style='List Paragraph')
    _set_bullet_numPr(p)
    return _apply_runs(p, parts)


def _code_block(doc, text):
    p = doc.add_paragraph(style='Normal')
    pf = p.paragraph_format
    pf.space_before = Pt(4)
    pf.space_after = Pt(4)
    r = p.add_run(text)
    r.font.name = 'Courier New'
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor.from_string(C_TEXT)
    return p


def _set_shading(cell, color):
    el = parse_xml('<w:shd {} w:fill="{}"/>'.format(nsdecls("w"), color))
    cell._tc.get_or_add_tcPr().append(el)


def _add_table(doc, headers, rows, caption=None):
    if caption:
        p = doc.add_paragraph(style='Caption')
        p.add_run(caption)

    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.alignment = 1
    tbl = t._tbl
    tblPr = tbl.tblPr if tbl.tblPr is not None else parse_xml(f'<w:tblPr {nsdecls("w")}/>')
    tblW = parse_xml(f'<w:tblW {nsdecls("w")} w:type="pct" w:w="5000"/>')
    existing_w = tblPr.find(qn('w:tblW'))
    if existing_w is not None:
        tblPr.remove(existing_w)
    tblPr.insert(0, tblW)
    tblLayout = tblPr.find(qn('w:tblLayout'))
    if tblLayout is not None:
        tblPr.remove(tblLayout)
    tblPr.append(parse_xml(f'<w:tblLayout {nsdecls("w")} w:type="fixed"/>'))

    # Column widths: headers must fit one line; distribute the rest by content length.
    num_cols = len(headers)
    TWIPS_PER_CHAR = 120
    CELL_PADDING = 200
    total_twips = 9072  # A4 portrait usable width

    col_min_width = []
    col_content_len = []
    for j in range(num_cols):
        h_clean = re.sub(r'\*\*(.+?)\*\*', r'\1', headers[j].strip())
        h_width = len(h_clean) * TWIPS_PER_CHAR + CELL_PADDING
        content_len = 0
        for row in rows:
            val = row[j] if j < len(row) else ''
            content_len = max(content_len, len(str(val).strip()))
        col_min_width.append(h_width)
        col_content_len.append(max(content_len, len(h_clean)))

    total_min = sum(col_min_width)
    if total_min <= total_twips:
        remaining = total_twips - total_min
        total_content = sum(col_content_len) or 1
        col_widths = [col_min_width[j] + int(remaining * col_content_len[j] / total_content)
                      for j in range(num_cols)]
    else:
        col_widths = [int(total_twips * col_min_width[j] / total_min) for j in range(num_cols)]

    diff = total_twips - sum(col_widths)
    if diff != 0 and col_widths:
        col_widths[-1] += diff

    tblGrid = tbl.find(qn('w:tblGrid'))
    if tblGrid is not None:
        tbl.remove(tblGrid)
    tblGrid = parse_xml(f'<w:tblGrid {nsdecls("w")}/>')
    for w in col_widths:
        tblGrid.append(parse_xml(f'<w:gridCol {nsdecls("w")} w:w="{w}"/>'))
    tbl.insert(1, tblGrid)

    for row_obj in t.rows:
        for j, cell in enumerate(row_obj.cells):
            if j < len(col_widths):
                cell.width = Twips(col_widths[j])

    for j, h in enumerate(headers):
        c = t.rows[0].cells[j]
        c.text = ''
        h_clean = re.sub(r'\*\*(.+?)\*\*', r'\1', h)
        r = c.paragraphs[0].add_run(h_clean)
        r.font.bold = True
        r.font.color.rgb = RGBColor.from_string(C_HEADER_TEXT)
        r.font.size = Pt(10)
        _set_shading(c, C_HEADER_FILL)
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
        c.paragraphs[0].paragraph_format.space_before = Pt(4)
        c.paragraphs[0].paragraph_format.space_after = Pt(4)
        c.vertical_alignment = 1

    for i, row in enumerate(rows):
        for j in range(len(headers)):
            c = t.rows[i + 1].cells[j]
            c.text = ''
            c.vertical_alignment = 1
            c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
            val = str(row[j] if j < len(row) else '')
            if has_inline_formatting(val):
                for item in parse_inline(val):
                    text, bold, italic = item[0], item[1], item[2]
                    is_code = item[3] if len(item) > 3 else False
                    r = c.paragraphs[0].add_run(text)
                    r.font.size = Pt(10)
                    r.font.color.rgb = RGBColor.from_string(C_TEXT)
                    if bold:
                        r.font.bold = True
                    if italic:
                        r.font.italic = True
                    if is_code:
                        r.font.name = 'Courier New'
                        r.font.size = Pt(9)
            else:
                r = c.paragraphs[0].add_run(val)
                r.font.size = Pt(10)
                r.font.color.rgb = RGBColor.from_string(C_TEXT)

    borders = (
        '<w:tblBorders {}>'
        '<w:top w:val="single" w:sz="4" w:space="0" w:color="auto"/>'
        '<w:left w:val="single" w:sz="4" w:space="0" w:color="auto"/>'
        '<w:bottom w:val="single" w:sz="4" w:space="0" w:color="auto"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="auto"/>'
        '<w:insideH w:val="single" w:sz="4" w:space="0" w:color="auto"/>'
        '<w:insideV w:val="single" w:sz="4" w:space="0" w:color="auto"/>'
        '</w:tblBorders>'
    ).format(nsdecls("w"))
    t._tbl.tblPr.append(parse_xml(borders))
    return t


def _set_header_footer(doc, doc_title, doc_code, version="v1.0"):
    section = doc.sections[-1]

    header = section.header
    header.is_linked_to_previous = False
    if header.paragraphs:
        hp = header.paragraphs[0]
        hp.clear()
        r1 = hp.add_run('VnResource Co., Ltd')
        r1.font.size = Pt(8)
        r1.font.color.rgb = RGBColor.from_string(C_MUTED)
        hp.add_run('\t')
        r2 = hp.add_run(doc_title.upper()[:50])
        r2.font.size = Pt(8)
        r2.font.bold = True
        r2.font.color.rgb = RGBColor.from_string(C_BRAND)
        hp.add_run('\t')
        r3 = hp.add_run('{} / {}'.format(doc_code, version))
        r3.font.size = Pt(8)
        r3.font.color.rgb = RGBColor.from_string(C_ACCENT)

    footer = section.footer
    footer.is_linked_to_previous = False
    if footer.paragraphs:
        fp = footer.paragraphs[0]
        fp.clear()
        r1 = fp.add_run('Lưu hành nội bộ | SDC - SA 05/{}'.format(version))
        r1.font.size = Pt(8)
        r1.font.color.rgb = RGBColor.from_string(C_MUTED)
        fp.add_run('\t')
        r2 = fp.add_run('Trang ')
        r2.font.size = Pt(8)
        r2.font.color.rgb = RGBColor.from_string(C_MUTED)


# ──────────────────────────────────────────────
# MARKDOWN PARSER
# ──────────────────────────────────────────────

def parse_inline(text):
    """Parse bold/italic/`code` -> list of (text, bold, italic, code) tuples."""
    parts = []
    for seg in re.split(r'(`[^`]+`)', text):
        if seg.startswith('`') and seg.endswith('`') and len(seg) > 2:
            parts.append((seg[1:-1], False, False, True))
        elif seg:
            pattern = re.compile(r'(\*\*\*(.+?)\*\*\*|\*\*(.+?)\*\*|\*(.+?)\*|([^*]+))')
            for m in pattern.finditer(seg):
                if m.group(2):
                    parts.append((m.group(2), True, True, False))
                elif m.group(3):
                    parts.append((m.group(3), True, False, False))
                elif m.group(4):
                    parts.append((m.group(4), False, True, False))
                elif m.group(5):
                    parts.append((m.group(5), False, False, False))
    if not parts:
        parts = [(text, False, False, False)]
    return parts


def has_inline_formatting(text):
    return bool(re.search(r'\*\*.*?\*\*|\*.*?\*|`.+?`', text))


def strip_md_links(text):
    return re.sub(r'\[([^\]]+)\]\([^)]+\)', r'\1', text)


def split_frontmatter(md_text):
    """Split a leading YAML `--- ... ---` frontmatter block. Returns (meta, body)."""
    meta = {}
    m = re.match(r'^﻿?---\s*\n(.*?)\n---\s*\n', md_text, re.DOTALL)
    if not m:
        return meta, md_text
    for line in m.group(1).splitlines():
        kv = re.match(r'\s*([A-Za-z0-9_\- ]+?)\s*:\s*(.*)', line)
        if kv:
            meta[kv.group(1).strip()] = kv.group(2).strip().strip('"\'')
    return meta, md_text[m.end():]


def parse_table(lines):
    headers, rows = [], []
    for i, line in enumerate(lines):
        cells = [c.strip() for c in line.strip().strip('|').split('|')]
        if i == 0:
            headers = cells
        elif i == 1:
            continue  # separator
        else:
            rows.append(cells)
    return headers, rows


def parse_md_to_blocks(md_text):
    lines = md_text.split('\n')
    blocks = []
    i = 0
    while i < len(lines):
        line = lines[i]

        if not line.strip():
            blocks.append(('empty', ''))
            i += 1
            continue

        if re.match(r'^---+\s*$', line.strip()):
            blocks.append(('hr', ''))
            i += 1
            continue

        m = re.match(r'^(#{1,4})\s+(.+)', line)
        if m:
            level = len(m.group(1))
            text = m.group(2).strip()
            # Strip manual numbering — Word auto-numbers via Heading styles.
            text = re.sub(r'^\d+(\.\d+)*\.?\s+', '', text)
            blocks.append(('h{}'.format(level), text))
            i += 1
            continue

        if line.strip().startswith('```'):
            code_lines = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith('```'):
                code_lines.append(lines[i])
                i += 1
            i += 1
            blocks.append(('code', '\n'.join(code_lines)))
            continue

        if line.strip().startswith('|'):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                table_lines.append(lines[i])
                i += 1
            blocks.append(('table', table_lines))
            continue

        if line.strip().startswith('>'):
            text = line.strip()[1:].strip()
            bq_bullet = re.match(r'^[-*]\s+(.+)', text)
            if bq_bullet:
                blocks.append(('bullet', strip_md_links(bq_bullet.group(1))))
            else:
                blocks.append(('blockquote', strip_md_links(text)))
            i += 1
            continue

        m = re.match(r'^(\s*)[-*]\s+(.+)', line)
        if m:
            blocks.append(('bullet', strip_md_links(m.group(2).strip())))
            i += 1
            continue

        m = re.match(r'^(\s*)\d+\.\s+(.+)', line)
        if m:
            blocks.append(('numbered', strip_md_links(m.group(2).strip())))
            i += 1
            continue

        blocks.append(('para', strip_md_links(line.strip())))
        i += 1

    return blocks


def extract_metadata(blocks):
    """Extract `> **Key:** value` metadata blockquotes at the top of the doc."""
    meta = {}
    for btype, bdata in blocks:
        if btype == 'blockquote':
            m = re.match(r'\*\*(.+?):\*\*\s*(.*)', bdata)
            if m:
                meta[m.group(1).strip()] = m.group(2).strip()
        elif btype in ('h1', 'h2', 'h3'):
            break
    return meta


def blocks_to_docx(doc, blocks):
    """Convert parsed blocks into DOCX content (main body)."""
    # H2-based file? (single title H1, ## as main sections) -> promote h2->H1 etc.
    content_h1_count = sum(1 for i, (bt, _) in enumerate(blocks) if bt == 'h1' and i > 0)
    promote = content_h1_count <= 1

    h1_count = 0
    table_count_per_h1 = 0

    skip_meta = True
    skip_toc = False
    skip_section_names = ('mục lục', 'muc luc', 'table of contents',
                          'danh sách bảng', 'danh sach bang')

    for idx, (btype, bdata) in enumerate(blocks):
        if btype == 'h1' and skip_meta:
            skip_meta = False
            continue
        if skip_meta and btype == 'blockquote':
            continue
        if skip_meta and btype in ('empty', 'hr'):
            continue
        if skip_meta and btype not in ('blockquote', 'empty', 'hr'):
            skip_meta = False

        if btype in ('h1', 'h2') and bdata.strip().lower() in skip_section_names:
            skip_toc = True
            continue
        if skip_toc:
            if btype in ('h1', 'h2') and bdata.strip().lower() not in skip_section_names:
                skip_toc = False
            else:
                continue

        if btype == 'empty' or btype == 'hr':
            continue

        elif btype == 'h1':
            h1_count += 1
            table_count_per_h1 = 0
            _h1(doc, bdata)

        elif btype == 'h2':
            if promote:
                h1_count += 1
                table_count_per_h1 = 0
                _h1(doc, bdata)
            else:
                _h2(doc, bdata)

        elif btype == 'h3':
            _h2(doc, bdata) if promote else _h3(doc, bdata)

        elif btype == 'h4':
            _h3(doc, bdata)

        elif btype == 'table':
            headers, rows = parse_table(bdata)
            cleaned_headers = [re.sub(r':\w+:', '', h).strip() for h in headers]
            cleaned_rows = [[re.sub(r':\w+:', '', strip_md_links(c)).strip() for c in row]
                            for row in rows]
            table_count_per_h1 += 1
            desc = ''
            for prev_idx in range(idx - 1, -1, -1):
                prev_type, prev_data = blocks[prev_idx]
                if prev_type == 'para':
                    desc = re.sub(r'\*{1,2}([^*]+)\*{1,2}', r'\1', prev_data.strip().rstrip(':'))
                    break
                elif prev_type in ('h1', 'h2', 'h3', 'h4'):
                    desc = prev_data.replace('`', '').strip()
                    break
                elif prev_type in ('empty', 'hr'):
                    continue
                else:
                    break
            n = max(h1_count, 1)
            caption = ("Bảng {}.{}. {}".format(n, table_count_per_h1, desc) if desc
                       else "Bảng {}.{}.".format(n, table_count_per_h1))
            _add_table(doc, cleaned_headers, cleaned_rows, caption=caption)

        elif btype in ('bullet', 'numbered'):
            _bullet_rich(doc, parse_inline(bdata))

        elif btype == 'code':
            _code_block(doc, bdata)

        elif btype == 'para':
            if bdata.startswith('*') and bdata.endswith('*') and not bdata.startswith('**'):
                p = _normal(doc, '')
                r = p.runs[0] if p.runs else p.add_run('')
                r.text = bdata.strip('*')
                r.font.italic = True
                continue
            parts = parse_inline(bdata)
            if len(parts) == 1 and parts[0][1]:
                _normal(doc, parts[0][0], bold=True)
            else:
                _normal_rich(doc, parts)

        elif btype == 'blockquote':
            p = doc.add_paragraph(style='Normal')
            for item in parse_inline(bdata):
                text, bold, italic = item[0], item[1], item[2]
                is_code = item[3] if len(item) > 3 else False
                r = p.add_run(text)
                r.font.italic = True
                r.font.color.rgb = RGBColor.from_string(C_MUTED)
                if bold:
                    r.font.bold = True
                if is_code:
                    r.font.name = 'Courier New'
                    r.font.size = Pt(9)


# ──────────────────────────────────────────────
# COVER PAGE + FRONT MATTER
# ──────────────────────────────────────────────

def _build_cover(doc, title, doc_code, version, project_code, cover_date, template_dir):
    logo_path = os.path.join(template_dir, 'logo-vnresource.png')
    if not os.path.exists(logo_path) and os.path.exists(DEFAULT_LOGO):
        logo_path = DEFAULT_LOGO
    if os.path.exists(logo_path):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(72)
        p.add_run().add_picture(logo_path, width=Cm(6))

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run('VnResource Co., Ltd')
    r.font.size = Pt(28)
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string(C_BRAND)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run('www.VnResource.vn')
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor.from_string(C_MUTED)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(20)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(title.upper())
    r.font.size = Pt(26)
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string(C_BRAND)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(24)

    for label, value, val_color in [
        ('Mã hiệu dự án: ', project_code, C_MUTED),
        ('Mã hiệu tài liệu: ', doc_code, C_ACCENT),
        ('Phiên bản tài liệu: ', version, C_ACCENT),
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(label)
        r.font.size = Pt(12)
        r.font.bold = True
        r.font.color.rgb = RGBColor.from_string(C_TEXT)
        r = p.add_run(value)
        r.font.size = Pt(12)
        r.font.color.rgb = RGBColor.from_string(val_color)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(216)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(cover_date)
    r.font.size = Pt(12)
    r.font.italic = True
    r.font.color.rgb = RGBColor.from_string(C_MUTED)

    _section_break(doc)


def _build_frontmatter(doc, version, update_date, author):
    # Revision history
    _h1_no_num(doc, 'Lịch sử cập nhật')
    _empty(doc)
    _add_table(doc,
               ['Phiên bản', 'Ngày cập nhật', 'Người cập nhật', 'Nội dung thay đổi'],
               [[version, update_date, author, 'Khởi tạo tài liệu']])
    _empty(doc)

    # Approval history
    _section_break(doc)
    _h1_no_num(doc, 'Lịch sử phê duyệt')
    _empty(doc)
    _add_table(doc,
               ['Người duyệt', 'Ngày duyệt', 'Nội dung duyệt', 'Ghi chú'],
               [['{Người duyệt}', '{Ngày duyệt}', '{Nội dung duyệt}', '{Ghi chú}']])
    _empty(doc)

    # Table of contents (auto-updating field)
    _section_break(doc)
    _h1_no_num(doc, 'Mục lục')
    p = doc.add_paragraph()
    r = p.add_run()
    r._r.append(parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="begin"/>'))
    r2 = p.add_run()
    r2._r.append(parse_xml(
        f'<w:instrText {nsdecls("w")} xml:space="preserve"> TOC \\o "1-3" \\h \\z \\u </w:instrText>'))
    r3 = p.add_run()
    r3._r.append(parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="separate"/>'))
    r4 = p.add_run('Nhấn Ctrl+A rồi F9 để cập nhật mục lục')
    r4.font.color.rgb = RGBColor.from_string('808080')
    r4.font.italic = True
    r5 = p.add_run()
    r5._r.append(parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="end"/>'))
    _section_break(doc)


# ──────────────────────────────────────────────
# ORCHESTRATION
# ──────────────────────────────────────────────

def _derive_code_from_filename(path):
    m = re.search(r'(VNR-[A-Z]+-\d+)', os.path.basename(path), re.IGNORECASE)
    return m.group(1).upper() if m else None


def convert(md_path, out_path, template, code=None, title=None, version=None,
            project_code='<Mã dự án>', cover_date=None, author=None,
            with_cover=True, quiet=False):
    with open(md_path, 'r', encoding='utf-8') as f:
        md_text = f.read()

    fm_meta, body = split_frontmatter(md_text)
    blocks = parse_md_to_blocks(body)
    bq_meta = extract_metadata(blocks)
    meta = {**fm_meta, **bq_meta}  # blockquote wins on conflict (in-body source of truth)

    # Resolve metadata with precedence: CLI arg > frontmatter/blockquote > default.
    if not code:
        code = (_derive_code_from_filename(md_path)
                or meta.get('Mã tài liệu') or meta.get('code') or 'VNR-XXX-000')
    if not version:
        version = meta.get('Phiên bản') or meta.get('Version') or meta.get('version') or 'v1.0'
    if not title:
        title = meta.get('title')
        if not title:
            for btype, bdata in blocks:
                if btype == 'h1':
                    title = re.sub(r'^(OVERVIEW|GUIDELINE|REFERENCE):\s*', '', bdata).strip()
                    break
        if not title:
            title = os.path.splitext(os.path.basename(md_path))[0]
    if not author:
        author = meta.get('Author') or meta.get('author') or 'SA Team'
    if not cover_date:
        now = datetime.date.today()
        cover_date = 'TPHCM, {:02d}/{}'.format(now.month, now.year)

    raw_update = (meta.get('Cập nhật lần cuối') or meta.get('Ngày tạo')
                  or datetime.date.today().isoformat())
    dm = re.search(r'\d{4}-\d{2}-\d{2}', raw_update)
    update_date = dm.group(0) if dm else datetime.date.today().isoformat()

    os.makedirs(os.path.dirname(out_path) or '.', exist_ok=True)
    shutil.copy2(template, out_path)
    doc = Document(out_path)
    clear_body(doc)

    if with_cover:
        _build_cover(doc, title, code, version, project_code, cover_date,
                     os.path.dirname(template))
        _build_frontmatter(doc, version, update_date, author)

    blocks_to_docx(doc, blocks)
    _set_header_footer(doc, title, code, version)

    doc.save(out_path)
    if not quiet:
        size = os.path.getsize(out_path)
        print("  [OK] {} -> {} ({} bytes)".format(code, os.path.basename(out_path), size))
    return out_path


def _resolve_output(md_path, output):
    md_stem = os.path.splitext(os.path.basename(md_path))[0]
    if not output:
        return os.path.join(os.path.dirname(os.path.abspath(md_path)), md_stem + '.docx')
    # Directory (existing dir, trailing sep, or no .docx extension) -> derive filename.
    is_dir = (os.path.isdir(output) or output.endswith(('/', '\\'))
              or not output.lower().endswith('.docx'))
    return os.path.join(output, md_stem + '.docx') if is_dir else output


def main():
    ap = argparse.ArgumentParser(
        description='Convert Markdown to a branded VNR DOCX using a template.')
    ap.add_argument('input', help='Input markdown file (.md)')
    ap.add_argument('-o', '--output', help='Output .docx file or directory')
    ap.add_argument('-t', '--template', default=DEFAULT_TEMPLATE, help='DOCX template to clone')
    ap.add_argument('--code', help='Document code, e.g. VNR-KH-001')
    ap.add_argument('--title', help='Cover title (default: first H1)')
    ap.add_argument('--version', help='Version, e.g. v1.0')
    ap.add_argument('--project-code', default='<Mã dự án>', help='Project code on cover')
    ap.add_argument('--date', dest='cover_date', help='Cover date, e.g. "TPHCM, 05/2026"')
    ap.add_argument('--author', help='Author for revision-history table')
    ap.add_argument('--no-cover', dest='with_cover', action='store_false',
                    help='Skip cover page + front matter')
    ap.add_argument('--quiet', action='store_true', help='Suppress log output')
    args = ap.parse_args()

    if not os.path.isfile(args.input):
        print('ERROR: input not found: {}'.format(args.input), file=sys.stderr)
        sys.exit(1)
    if not os.path.isfile(args.template):
        print('ERROR: template not found: {}'.format(args.template), file=sys.stderr)
        sys.exit(1)

    out_path = _resolve_output(args.input, args.output)
    try:
        convert(args.input, out_path, args.template,
                code=args.code, title=args.title, version=args.version,
                project_code=args.project_code, cover_date=args.cover_date,
                author=args.author, with_cover=args.with_cover, quiet=args.quiet)
    except Exception as e:
        print('  [FAIL] {} - {}'.format(args.input, e), file=sys.stderr)
        import traceback
        traceback.print_exc()
        sys.exit(1)


if __name__ == '__main__':
    main()
